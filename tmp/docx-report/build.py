import json,re
from pathlib import Path
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
source=Path(r'C:/Users/DatPNT/.codex/sessions/2026/10/16/rollout-2026-10-16T22-23-32-01a1454f-b006-7241-b6ad-e5e4d452724c.jsonl')
candidates=[]
for line in source.open(encoding='utf-8'):
 try: obj=json.loads(line)
 except: continue
 p=obj.get('payload',{})
 if p.get('role')=='assistant':
  txt='\n'.join(x.get('text','') for x in p.get('content',[]) if isinstance(x,dict))
  if '# CHƯƠNG 2.' in txt and '# CHƯƠNG 4.' in txt: candidates.append(txt)
assert candidates, 'Report not found in current conversation'
text=candidates[-1]
text=text[text.index('# CHƯƠNG 2.'):]
Path('tmp/docx-report/source.md').write_text(text,encoding='utf-8')
doc=Document(); sec=doc.sections[0]
sec.page_height=Cm(29.7);sec.page_width=Cm(21)
sec.top_margin=Cm(2.2);sec.bottom_margin=Cm(2.2);sec.left_margin=Cm(2.7);sec.right_margin=Cm(2)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet','List Number']:
 st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
 st.font.size=Pt(12)
 st.paragraph_format.space_after=Pt(6)
 st.paragraph_format.line_spacing=1.15
for name,size in [('Title',22),('Heading 1',17),('Heading 2',14),('Heading 3',12)]:
 st=doc.styles[name];st.font.size=Pt(size);st.font.bold=True
 st.paragraph_format.keep_with_next=True
 st.paragraph_format.space_before=Pt(12)
 if name=='Heading 1': st.paragraph_format.page_break_before=True
p=doc.add_paragraph('BÁO CÁO ĐỒ ÁN',style='Title');p.alignment=1
p=doc.add_paragraph('HÀNH TRÌNH AN NHIÊN',style='Title');p.alignment=1
p=doc.add_paragraph('Xây dựng ứng dụng di động hỗ trợ hành hương và tụng kinh',style='Subtitle');p.alignment=1
p=doc.add_paragraph('Chương 2 Cơ sở lý thuyết và công nghệ sử dụng\nChương 3 Phân tích và thiết kế hệ thống\nChương 4 Hiện thực hệ thống kiểm thử và triển khai');p.alignment=1
p=doc.add_paragraph('Digital Pilgrimage App\nTháng 10 năm 2026');p.alignment=1
footer=sec.footer.paragraphs[0];footer.alignment=1
r=footer.add_run();f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');r._r.addnext(f)
def inline(p,s):
 for chunk in re.split(r'(\*\*.*?\*\*|`[^`]+`)',s):
  if not chunk:continue
  run=p.add_run(chunk[2:-2] if chunk.startswith('**') else chunk[1:-1] if chunk.startswith('`') else chunk)
  if chunk.startswith('**'):run.bold=True
  if chunk.startswith('`'):run.font.name='Consolas';run.font.size=Pt(10)
lines=text.splitlines();i=0
while i<len(lines):
 s=lines[i].strip();i+=1
 if not s or s=='---':continue
 if s.startswith('```'):
  code=[]
  while i<len(lines) and not lines[i].startswith('```'):code.append(lines[i]);i+=1
  i+=1
  p=doc.add_paragraph();p.paragraph_format.line_spacing=1;p.paragraph_format.space_before=Pt(5)
  p.paragraph_format.keep_together=True
  r=p.add_run('\n'.join(code));r.font.name='Consolas';r.font.size=Pt(9)
  continue
 if s.startswith('|'):
  rows=[s]
  while i<len(lines) and lines[i].strip().startswith('|'):rows.append(lines[i].strip());i+=1
  rows=[x for x in rows if not re.match(r'^\|[\s:|\-]+$',x)]
  data=[[v.strip() for v in x.strip('|').split('|')] for x in rows]
  table=doc.add_table(rows=0,cols=len(data[0]));table.autofit=False
  widths={2:[5,11.3],3:[4,4.8,7.5],4:[3.4,4.3,3.4,5.2]}.get(len(data[0]),[16.3/len(data[0])]*len(data[0]))
  for col,w in zip(table.columns,widths):col.width=Cm(w)
  borders=OxmlElement('w:tblBorders')
  for edge in ['top','left','bottom','right','insideH','insideV']:
   e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
  table._tbl.tblPr.append(borders)
  for n,row in enumerate(data):
   cells=table.add_row().cells
   trPr=table.rows[-1]._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
   if n==0:trPr.append(OxmlElement('w:tblHeader'))
   for j,val in enumerate(row):
    c=cells[j];c.width=Cm(widths[j]);c.vertical_alignment=1
    p=c.paragraphs[0];p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4);p.paragraph_format.line_spacing=1.05
    inline(p,val)
    for r in p.runs:r.font.size=Pt(10);r.bold=n==0
    if n==0:
     sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EDF2');c._tc.get_or_add_tcPr().append(sh)
  doc.add_paragraph().paragraph_format.space_after=Pt(0)
  continue
 m=re.match(r'^(#{1,3})\s+(.*)',s)
 if m:
  p=doc.add_paragraph(m[2],style='Heading '+str(len(m[1])));continue
 style=None
 if s.startswith('- '):style='List Bullet';s=s[2:]
 elif re.match(r'^\d+\.\s',s):style=None
 p=doc.add_paragraph(style=style);inline(p,s)
 if not style:p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 p.paragraph_format.widow_control=True
 if re.match(r'^\d+\.\s',s):
  p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.05
out=Path('Bao_cao_Chuong_2_3_4_Hanh_Trinh_An_Nhien.docx')
for st in doc.styles:
 for border in list(st.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
for p in doc.paragraphs:
 for border in list(p._p.iter(qn('w:pBdr'))): border.getparent().remove(border)
doc.save(out)
print('Saved:',out.resolve());print('Source characters:',len(text));print('Tables:',len(doc.tables));print('Paragraphs:',len(doc.paragraphs))
